"""Generate a NexusNet pro se patent drafting packet.

This script intentionally writes all generated artifacts under the repository
root. It does not write NexusNet packet state to user-home caches.
"""

from __future__ import annotations

import datetime as dt
import textwrap
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch
from reportlab.pdfgen import canvas


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "patent" / "nexusnet_micro_entity_packet_2026-05-01"
DOCX_DIR = OUT / "docx"
DRAWINGS_DIR = OUT / "drawings"
LETTERS_DIR = OUT / "letters"
QA_DIR = OUT / "_qa_renders"

DATE = "2026-05-01"
TITLE = "Hive Mind Harness for Artificial Intelligence with Governed Sparse Routing"

USPTO_SOURCES = [
    ("USPTO nonprovisional utility applications", "https://www.uspto.gov/patents-getting-started/patent-basics/types-patent-applications/nonprovisional-utility-patent"),
    ("USPTO provisional applications", "https://www.uspto.gov/patents-getting-started/patent-basics/types-patent-applications/provisional-application-patent"),
    ("USPTO patent forms", "https://www.uspto.gov/patents/apply/forms"),
    ("USPTO micro entity status", "https://www.uspto.gov/patents/laws/micro-entity-status"),
    ("USPTO fee schedule", "https://www.uspto.gov/learning-and-resources/fees-and-payment/uspto-fee-schedule"),
    ("USPTO DOCX filing information", "https://www.uspto.gov/patents/docx"),
    ("USPTO Patent Center", "https://patentcenter.uspto.gov/"),
    ("USPTO patent drawing rules overview", "https://www.uspto.gov/patents/apply/applying-online/technical-requirements-documents-submitted-uspto"),
]


def wrap(text: str) -> str:
    return textwrap.dedent(text).strip() + "\n"


def write(path: Path, content: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content, encoding="utf-8", newline="\n")


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_para(doc: Document, text: str = "", style: str | None = None) -> None:
    p = doc.add_paragraph(style=style)
    for chunk in text.split("\n"):
        if chunk:
            p.add_run(chunk)
        p.add_run().add_break() if chunk != text.split("\n")[-1] else None


def add_bullet(doc: Document, text: str) -> None:
    doc.add_paragraph(text, style="List Bullet")


def style_doc(doc: Document) -> None:
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)
    for name, size, color in [
        ("Title", 18, RGBColor(31, 78, 121)),
        ("Heading 1", 14, RGBColor(31, 78, 121)),
        ("Heading 2", 12, RGBColor(47, 84, 150)),
        ("Heading 3", 11, RGBColor(68, 68, 68)),
    ]:
        styles[name].font.name = "Arial"
        styles[name].font.size = Pt(size)
        styles[name].font.color.rgb = color
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)


def source_links_md() -> str:
    return "\n".join(f"- [{label}]({url})" for label, url in USPTO_SOURCES)


README = wrap(f"""
    # NexusNet Pro Se Patent Drafting Packet

    Generated: {DATE}

    Working title: **{TITLE}**

    This packet is a drafting aid for a United States pro se utility patent
    filing by a potential micro-entity applicant. It is not legal advice, does
    not determine patentability, and does not replace review by a registered
    patent practitioner. Software and AI inventions often need careful claim
    drafting to address subject-matter eligibility, novelty, obviousness,
    enablement, written description, inventorship, ownership, and disclosure
    timing.

    ## What This Packet Contains

    - A provisional/nonprovisional filing roadmap.
    - A micro-entity eligibility worksheet.
    - USPTO form field worksheets for ADS, declaration, provisional cover
      sheet, micro-entity certification, and fee/transmittal steps.
    - A working invention disclosure record.
    - A nonprovisional-style specification draft.
    - A claim set draft with system, method, and computer-readable medium claim
      families.
    - An abstract draft.
    - Formal drawing source files and a drawing PDF.
    - Prior-art and differentiator worksheets.
    - Transmittal and cover-note templates.

    ## Generated Files

    Use the Markdown files as the editable source. Use the DOCX files as working
    Patent Center drafting material. Use the drawings PDF as a starting drawing
    packet, then check it against current USPTO drawing requirements before
    filing.

    ## Official Sources Checked

    {source_links_md()}

    ## Immediate Pro Se Workflow

    1. Complete `01_MICRO_ENTITY_WORKSHEET.md`.
    2. Complete personal identity, inventor, residence, correspondence,
       citizenship, and ownership fields in `02_USPTO_FORM_FIELD_WORKSHEETS.md`.
    3. Review `03_INVENTION_DISCLOSURE_RECORD.md` and add any missing
       implementation details, dates, public disclosures, contributors, or
       ownership facts.
    4. Decide whether to file a provisional first or a nonprovisional utility
       application now.
    5. If filing nonprovisional, review `04_NONPROVISIONAL_SPECIFICATION_DRAFT.md`,
       `05_CLAIMS_DRAFT.md`, `06_ABSTRACT_DRAFT.md`, and the drawings.
    6. Download current official USPTO forms from the USPTO forms page and
       transfer the worksheet fields into the official PDFs or Patent Center
       screens.
    7. File through Patent Center using the official current fee schedule and
       micro-entity certification if you qualify.

    ## Important Disclosure Warning

    If NexusNet has already been publicly disclosed, offered for sale, sold,
    published, posted, demonstrated, or shared without confidentiality, filing
    deadlines and foreign rights may be affected. Record every such disclosure
    in `03_INVENTION_DISCLOSURE_RECORD.md` before filing.
""")


USE_LIMITS = wrap("""
    # Use Limits And Legal Caution

    This packet is generated from the current NexusNet source canon and project
    architecture. It is a drafting aid, not legal advice.

    ## You Must Personally Verify

    - Correct inventorship.
    - Whether anyone else contributed to claimed subject matter.
    - Whether any employer, contractor, investor, or collaborator owns rights.
    - Whether any public disclosure already occurred.
    - Whether micro-entity status applies to every applicant/inventor.
    - Whether the claims are patent-eligible, novel, and non-obvious.
    - Whether the specification enables the full claimed scope.
    - Whether any open-source licenses affect what can be claimed or commercialized.

    ## Recommended Minimum Professional Review

    Even if you file pro se, consider a limited-scope review by a registered
    patent attorney or patent agent before filing claims. Claim drafting is the
    part most likely to determine whether the filing has commercial value.
""")


READ_ORDER = wrap("""
    # Read Me First - Patent Packet Reading Order

    This is the one document to use as your map. The packet is not meant to be
    read strictly from file `00` to file `12` without context. Some files are
    mandatory first-pass reading, some are decision points, and some are
    reference material you only need when you reach that filing step.

    ## Short Version

    Read these in this order:

    1. `00_USE_AND_LIMITS.md`
    2. `README_PRO_SE_PACKET.md`
    3. `03_INVENTION_DISCLOSURE_RECORD.md`
    4. `01_MICRO_ENTITY_WORKSHEET.md`
    5. `09_FILING_CHECKLIST_AND_PATENT_CENTER_STEPS.md`
    6. Choose one path:
       - Provisional path: `12_PROVISIONAL_APPLICATION_DISCLOSURE_DRAFT.md`
       - Nonprovisional path: `04_NONPROVISIONAL_SPECIFICATION_DRAFT.md`,
         then `05_CLAIMS_DRAFT.md`, then `06_ABSTRACT_DRAFT.md`
    7. `07_DRAWING_BRIEF_DESCRIPTION_AND_REFERENCE_NUMERALS.md`
    8. `11_SOURCE_CANON_TO_CLAIM_SUPPORT_CHART.md`
    9. `08_PRIOR_ART_AND_DIFFERENTIATOR_WORKSHEET.md`
    10. `10_IDS_WORKSHEET.md`
    11. `02_USPTO_FORM_FIELD_WORKSHEETS.md`
    12. The letter templates in `letters/`

    ## Why This Order

    ### 1. Read The Warning First

    Start with `00_USE_AND_LIMITS.md`.

    This tells you what the packet is and is not. It also flags the things I
    cannot know for you: inventorship, ownership, disclosure timing,
    micro-entity eligibility, employer rights, assignments, patentability, and
    whether a practitioner should review the claims.

    Do not skip this one.

    ### 2. Read The Packet Overview

    Next read `README_PRO_SE_PACKET.md`.

    This explains what each artifact is for and links the official USPTO
    sources used when building the packet.

    ### 3. Fill The Invention Disclosure Record Before Filing Anything

    Read and fill `03_INVENTION_DISCLOSURE_RECORD.md` early.

    This is where you record:

    - what the invention is
    - what concepts you believe are inventive
    - who contributed
    - whether anything was publicly disclosed
    - whether any employment, contractor, investor, or open-source issue affects
      ownership
    - whether anything has been sold, offered, demoed, posted, or shared

    This file determines whether the rest of the packet is safe to rely on.

    ### 4. Check Micro-Entity Eligibility Early

    Read `01_MICRO_ENTITY_WORKSHEET.md`.

    This tells you what to verify before claiming micro-entity status. If you
    are not sure, do not guess. Use the current USPTO micro-entity page and
    forms.

    ### 5. Decide Filing Path

    Read `09_FILING_CHECKLIST_AND_PATENT_CENTER_STEPS.md`.

    This is where you choose whether you are trying to file:

    - a provisional application first, or
    - a nonprovisional utility application now.

    If you are unsure, the provisional path is often the easier first pro se
    drafting path, but it does not become a patent by itself and must be
    followed by a nonprovisional within the applicable deadline.

    ## If You Choose The Provisional Path

    Read:

    1. `12_PROVISIONAL_APPLICATION_DISCLOSURE_DRAFT.md`
    2. `07_DRAWING_BRIEF_DESCRIPTION_AND_REFERENCE_NUMERALS.md`
    3. `drawings/NEXUSNET_PATENT_DRAWINGS.pdf`
    4. `02_USPTO_FORM_FIELD_WORKSHEETS.md`, only the provisional cover sheet
       and micro-entity parts

    The provisional path is more about preserving a disclosure and filing date
    than perfecting claims. Still, the disclosure needs enough technical detail
    to support whatever you may later claim.

    ## If You Choose The Nonprovisional Path

    Read:

    1. `04_NONPROVISIONAL_SPECIFICATION_DRAFT.md`
    2. `05_CLAIMS_DRAFT.md`
    3. `06_ABSTRACT_DRAFT.md`
    4. `07_DRAWING_BRIEF_DESCRIPTION_AND_REFERENCE_NUMERALS.md`
    5. `drawings/NEXUSNET_PATENT_DRAWINGS.pdf`
    6. `11_SOURCE_CANON_TO_CLAIM_SUPPORT_CHART.md`

    Then use:

    - `docx/NEXUSNET_UTILITY_SPECIFICATION_CLAIMS_ABSTRACT_DRAFT.docx`
    - official USPTO forms or Patent Center screens

    The claim draft is the most legally sensitive part of the packet. Treat it
    as a starting point, not a finished legal product.

    ## Read These Before Final Submission

    Read `08_PRIOR_ART_AND_DIFFERENTIATOR_WORKSHEET.md` before filing claims.

    This helps you think through what already exists and how NexusNet is
    different. It is also useful when deciding whether the claims are too broad,
    too vague, or missing the real technical improvement.

    Read `10_IDS_WORKSHEET.md` before or shortly after filing.

    This is where you track references that may need to be disclosed to the
    USPTO. It is not an official IDS form.

    Read `02_USPTO_FORM_FIELD_WORKSHEETS.md` when you are ready to complete
    official forms or Patent Center screens.

    ## Files You Can Read Whenever

    These are reference/helper files. You do not need to read them in order:

    - `PACKAGE_MANIFEST.md`
    - `drawings/*.svg`
    - `letters/TRANSMITTAL_LETTER_TEMPLATE.md`
    - `letters/MICRO_ENTITY_COVER_NOTE.md`
    - `letters/INVENTOR_DECLARATION_PREP_NOTE.md`
    - `docx/NEXUSNET_PRO_SE_FORMS_AND_WORKSHEETS.docx`

    ## Practical First Sitting

    If you only have one focused session, do this:

    1. Read `00_USE_AND_LIMITS.md`.
    2. Read `README_PRO_SE_PACKET.md`.
    3. Fill as much as possible in `03_INVENTION_DISCLOSURE_RECORD.md`.
    4. Fill `01_MICRO_ENTITY_WORKSHEET.md`.
    5. Read `09_FILING_CHECKLIST_AND_PATENT_CENTER_STEPS.md`.
    6. Decide provisional versus nonprovisional.

    Stop there if you are tired. That decision controls which drafting path you
    read next.
""")


MICRO_ENTITY = wrap("""
    # Micro-Entity Eligibility Worksheet

    Use this worksheet before claiming micro-entity status. Transfer answers
    into the current official USPTO micro-entity certification form only after
    confirming the facts.

    Official source: https://www.uspto.gov/patents/laws/micro-entity-status

    ## Route A: Gross Income Basis

    The usual form is PTO/SB/15A or its current equivalent. Confirm current
    form numbers on the USPTO forms page before filing.

    1. Applicant qualifies as a small entity:
       - [ ] Yes
       - [ ] No
       - Notes:

    2. Each inventor/applicant has not been named as inventor on more than four
       prior U.S. nonprovisional patent applications, excluding allowed
       exceptions:
       - [ ] Yes
       - [ ] No
       - Prior applications to review:

    3. Each inventor/applicant's gross income for the required prior calendar
       year is below the current USPTO micro-entity income threshold:
       - [ ] Yes
       - [ ] No
       - Threshold checked on:
       - Income year checked:

    4. No inventor/applicant has assigned, granted, or conveyed, and is not
       obligated to assign, grant, or convey, rights to an entity exceeding the
       income threshold:
       - [ ] Yes
       - [ ] No
       - Assignment/obligation notes:

    5. All inventors/applicants can sign or otherwise support the certification:
       - [ ] Yes
       - [ ] No

    ## Route B: Institution Of Higher Education Basis

    The usual form is PTO/SB/15B or its current equivalent. Use only if the
    higher-education requirements actually apply.

    - [ ] Applicant is employed by an institution of higher education.
    - [ ] Applicant has assigned or is obligated to assign rights to such an
      institution.
    - [ ] Institution details:

    ## Maintenance Warning

    Micro-entity status can change. Re-check before paying later fees.
""")


FORMS = wrap("""
    # USPTO Form Field Worksheets

    These worksheets help you transfer information into current official USPTO
    forms or Patent Center screens. They are not substitutes for the official
    USPTO forms.

    Official forms page: https://www.uspto.gov/patents/apply/forms

    ## Application Data Sheet Worksheet

    Common form: PTO/AIA/14 or current Application Data Sheet equivalent.

    - Application type:
      - [ ] Provisional
      - [ ] Nonprovisional utility
    - Invention title:
      Hive Mind Harness for Artificial Intelligence with Governed Sparse Routing
    - Inventor legal name:
      [INSERT FULL LEGAL NAME]
    - Inventor residence city/state/country:
      [INSERT]
    - Inventor mailing address:
      [INSERT]
    - Inventor citizenship:
      [INSERT]
    - Correspondence address:
      [INSERT]
    - Applicant:
      - [ ] Inventor
      - [ ] Assignee
      - [ ] Other
    - Assignee or ownership interest, if any:
      [INSERT OR NONE]
    - Domestic benefit or priority claim:
      [INSERT IF CLAIMING PRIOR PROVISIONAL/NONPROVISIONAL]
    - Foreign priority claim:
      [INSERT IF ANY]
    - Attorney docket number:
      NEXUSNET-001-US

    ## Inventor Oath Or Declaration Worksheet

    Common form: PTO/AIA/01 or current declaration equivalent.

    - Inventor legal name:
      [INSERT]
    - Application title:
      Hive Mind Harness for Artificial Intelligence with Governed Sparse Routing
    - Inventor confirms original inventor status:
      - [ ] Yes
    - Inventor acknowledges duty to disclose material information:
      - [ ] Yes
    - Signature:
      [SIGN ONLY ON OFFICIAL FORM OR PATENT CENTER EQUIVALENT]
    - Date:
      [INSERT]

    ## Micro-Entity Certification Worksheet

    Common form: PTO/SB/15A for gross income basis.

    - Basis:
      - [ ] Gross income
      - [ ] Institution of higher education
    - Small entity confirmed:
      - [ ] Yes
    - Prior application limit confirmed:
      - [ ] Yes
    - Income threshold confirmed:
      - [ ] Yes
    - Assignment/obligation threshold confirmed:
      - [ ] Yes
    - Signature:
      [SIGN ONLY ON OFFICIAL FORM OR PATENT CENTER EQUIVALENT]

    ## Provisional Cover Sheet Worksheet

    Common form: PTO/SB/16 or current provisional cover sheet equivalent.

    - Title:
      Hive Mind Harness for Artificial Intelligence with Governed Sparse Routing
    - Inventor:
      [INSERT]
    - Correspondence address:
      [INSERT]
    - Entity status:
      - [ ] Micro entity
      - [ ] Small entity
      - [ ] Undiscounted
    - Government support:
      [NONE OR INSERT]
    - Attorney docket number:
      NEXUSNET-001-PROV

    ## Fee Transmittal / Patent Center Fee Worksheet

    Check the current USPTO fee schedule before paying.

    - Filing fee:
      [CHECK CURRENT MICRO-ENTITY AMOUNT]
    - Search fee, if nonprovisional:
      [CHECK CURRENT MICRO-ENTITY AMOUNT]
    - Examination fee, if nonprovisional:
      [CHECK CURRENT MICRO-ENTITY AMOUNT]
    - Excess claims fees:
      [CHECK AFTER FINAL CLAIM COUNT]
    - Application size fee:
      [CHECK AFTER FINAL PAGE COUNT]
""")


DISCLOSURE = wrap("""
    # NexusNet Invention Disclosure Record

    ## Working Title

    Hive Mind Harness for Artificial Intelligence with Governed Sparse Routing

    ## Short Description

    NexusNet is a brain-first artificial intelligence architecture in which a
    canonical neural core mediates all task execution through a modular hive of
    assistant orchestrators, expert capsules, memory planes, sandboxed tools,
    evaluators, checkpoint ledgers, and privacy-preserving federated learning
    hooks. The system treats operational software components as neural-harness
    nodes, routes typed activations through a neural bus, activates sparse
    experts through a cortex router, performs recurrent deliberation before
    action, and permits self-improvement only through sandbox, evaluation,
    immune/governance, and rollback gates.

    ## Inventive Concepts To Preserve

    1. A brain-mediated AI harness in which user tasks, tools, models,
       providers, skill systems, and agent bridges cannot bypass a canonical
       neural core.
    2. A software neural substrate mapping nodes, edges, weights, activations,
       forward passes, recurrent loops, loss, optimization, and checkpoints to
       live AI harness components.
    3. A sparse MoE-style router that selects assistant orchestrators, expert
       capsules, Mini-NexusNets, tools, memory banks, sandboxes, providers, and
       evaluators based on confidence, risk, capability, privacy, cost, latency,
       and historical reliability.
    4. A multi-plane memory and engram layer that distinguishes canon, addenda,
       traces, evidence, temporal truth, dream artifacts, training artifacts,
       and federated lessons.
    5. A self-assimilation loop in which new candidate capabilities are
       researched, source-pinned, licensed, privacy-classified, sandbox-tested,
       evaluated, promoted, rejected, blocked, or side-barred.
    6. A checkpoint/rewind ledger requiring pre-write snapshots and rollback
       metadata for system evolution.
    7. A federated learning plane that shares redacted route scores, expert
       deltas, runtime scorecards, and evaluation summaries without raw private
       data.
    8. A VisualOps cockpit that renders brain path traces, AO activity, expert
       activation, memory use, security gates, route disagreements, training
       state, federation state, and product readiness.

    ## Known Source Canon Inputs

    - docs/NEXUSNET_COMPLETE_CHAT_CANON_BOOK_2026-04-28.md
    - docs/NEXUSNET_CHAT_RESEARCH_CANON_2026-04-28.md
    - docs/NEXUSNET_POST_BOOK_CANON_ADDENDUM_2026-05-01.md
    - docs/superpowers/specs/2026-05-01-nexusnet-hive-neural-substrate-v0-design.md
    - nexusnet/core/brain.py
    - nexusnet/canon/realization.py

    ## Public Disclosure Log

    Complete this before filing.

    | Date | Disclosure type | Audience | Confidential? | Evidence/location |
    | --- | --- | --- | --- | --- |
    | [INSERT] | [GitHub, demo, chat export, video, sale, offer, pitch, etc.] | [INSERT] | [YES/NO] | [INSERT] |

    ## Contributor / Inventor Review

    List every person who contributed to the claimed technical concepts, not
    merely people who followed instructions or did ordinary implementation.

    | Person | Contribution | Claimed concept affected | Employment/contract status | Assignment status |
    | --- | --- | --- | --- | --- |
    | [INSERT] | [INSERT] | [INSERT] | [INSERT] | [INSERT] |

    ## Ownership Review

    - Employer rights:
    - Contractor rights:
    - Investor rights:
    - Open-source dependency constraints:
    - Prior assignment obligations:
    - Government funding:

    ## Commercial Use / Sale Review

    - Has the invention been offered for sale?
    - Has a buyer received a build or executable?
    - Has a public repository contained enabling details?
    - Has any private disclosure lacked an NDA?
""")


SPEC = wrap("""
    # Nonprovisional Utility Specification Draft

    ## Title

    Hive Mind Harness for Artificial Intelligence with Governed Sparse Routing

    ## Cross-Reference To Related Applications

    [If filing after a provisional application, insert the provisional
    application number and filing date here. Otherwise state: Not applicable.]

    ## Field

    The disclosure relates to artificial intelligence systems, multi-agent
    orchestration, machine-learning model routing, memory-augmented computing,
    software safety systems, federated learning, and computer-implemented
    training and evaluation harnesses.

    ## Background

    Conventional AI applications commonly route a user prompt directly to a
    hosted or local model, retrieve documents through a retrieval-augmented
    generation pipeline, or invoke tools through an agent framework. Such
    systems often treat memory, tool permissions, provider routing, model
    selection, evaluation, autonomous coding, and user-interface status as
    separate features. As a result, the systems can lack a unified decision
    authority, can fail to record why a route was selected, can permit unsafe
    tool use, can accumulate stale skills, and can treat self-improvement as an
    ad hoc update rather than as a reversible, evaluated, and policy-gated
    process.

    Multi-agent systems can improve specialization, but many such systems
    activate broad agent sets, depend on prompt-only conventions, or lack a
    neural-network-like substrate for nodes, edges, activations, recurrent
    loops, loss signals, and optimization. Mixture-of-experts model
    architectures provide sparse expert activation at the model-weight level,
    but do not by themselves govern external tools, memory, sandboxed updates,
    provenance ledgers, visual operator controls, or privacy-preserving
    deployment-level learning.

    A need therefore exists for a computer-implemented AI harness that combines
    a brain-first command path, sparse expert activation, explicit memory
    planes, recurrent deliberation, sandboxed self-assimilation, trace-first
    evaluation, checkpoint-based rollback, and federated learning boundaries in
    a single governed architecture.

    ## Summary

    In some embodiments, a computer-implemented neural network harness includes
    a neural core service, a node registry, a neural bus, a cortex router,
    assistant orchestrators, expert capsules, a multi-plane memory system, a
    recurrent deliberation loop, an immune/governance kernel, a sandbox and
    evaluation subsystem, a checkpoint/rewind ledger, and an action/output
    subsystem. The harness receives an operator task or candidate capability,
    converts it into a typed activation, retrieves relevant memory and
    provenance, selects a sparse set of nodes, performs one or more internal
    deliberation loops, evaluates confidence and risk, and permits an external
    action only after required gates pass.

    In some embodiments, candidate improvements are not directly installed.
    Instead, a candidate is source-pinned, license-checked, privacy-classified,
    converted into a capability genome, tested in a closed sandbox, compared
    against baseline behavior, evaluated for security and regressions, and then
    promoted, rejected, blocked, or side-barred with ledger evidence.

    In some embodiments, the harness includes a federated learning plane that
    exports only approved, redacted deltas, such as route-score summaries,
    expert-performance scorecards, evaluation results, runtime-performance
    summaries, and failure signatures. Raw private prompts, secrets, local file
    contents, proprietary code, and unredacted transcripts are not exported
    without explicit approval.

    ## Brief Description Of The Drawings

    FIG. 1 illustrates an example brain-first neural network harness system.
    FIG. 2 illustrates example neural-harness planes.
    FIG. 3 illustrates an example task forward pass through the harness.
    FIG. 4 illustrates an example sandboxed self-assimilation loop.
    FIG. 5 illustrates an example skill-system and AFK sandbox agent factory.
    FIG. 6 illustrates an example privacy-preserving federated learning flow.
    FIG. 7 illustrates example trace, checkpoint, and rewind data structures.

    ## Detailed Description

    ### Definitions

    A "neural core" refers to a computer service that remains the authority for
    task routing, memory assembly, policy application, expert selection,
    generation, critique, evaluation, and final output metadata.

    A "hive node" refers to a typed harness component, including an assistant
    orchestrator, expert capsule, Mini-NexusNet, skill, skill system, tool
    adapter, model adapter, memory bank, evaluator, sandbox runner, policy gate,
    curator, or training school component.

    A "hive activation" refers to a typed representation of an input task,
    candidate update, event, tool output, or federated signal. The activation
    can include an intent, capability vector, risk vector, memory references,
    provenance references, confidence, novelty, privacy class, and policy
    labels.

    A "skill system" refers to an orchestrated workflow made of multiple
    focused reusable skills with typed handoffs, context limits, checkpoint
    locations, and visible output artifacts.

    ### System Overview

    Referring to FIG. 1, an operator input 102 enters a sensory/input plane 104.
    A representation plane 106 converts the input into a hive activation. A
    temporal lineage ledger 108 attaches session, dependency, prompt-preview,
    candidate-strain, and checkpoint lineage. A neural bus 110 publishes the
    activation to a cortex router 112 and a neural core 114. The cortex router
    selects a sparse subset of assistant orchestrators 116, expert capsules
    118, Mini-NexusNets 120, memory planes 122, sandbox/evaluation services 130,
    and provider/tool adapters 144.

    The neural core 114 mediates all external effects. A HiveBlackboard 124
    maintains residual working state across loops. A recurrent deliberation
    loop 126 repeats routing, memory lookup, expert computation, critique, and
    policy evaluation until an exit condition is satisfied. An immune/governance
    plane 128 blocks unsafe routes, prompt-injection attempts, untrusted tools,
    unapproved writes, unredacted federation, and candidate promotions lacking
    evidence. A checkpoint/rewind ledger 132 records reversible state before
    write actions. An action/output plane 134 emits approved code edits,
    reports, tool calls, API responses, generated assets, or visual state. A
    VisualOps control panel 136 renders traces and status.

    ### Neural-Harness Planes

    Referring to FIG. 2, the harness includes multiple planes corresponding to
    neural-network concepts. Nodes correspond to operational components, edges
    correspond to trust, dependency, provenance, and routing relationships,
    weights correspond to scores such as confidence, risk, usefulness, latency,
    privacy, cost, approval, and reliability, and activations correspond to
    typed task signals. A forward pass corresponds to a complete task execution.
    Recurrent loops correspond to repeated internal deliberation before external
    action. Loss corresponds to tests, regressions, policy violations, user
    corrections, cost drift, latency drift, unsupported claims, and usefulness
    deltas. Optimization is performed by a training school, curator, route
    score updates, evaluation gates, and federated aggregation.

    ### Sparse Expert Routing

    The cortex router 112 receives a focused activation set and scores candidate
    nodes. The score can include capability fit, historical reliability,
    confidence, privacy sensitivity, cost, latency, concurrency safety, tool
    permissions, certification state, quarantine state, memory access, and
    policy labels. The router selects only the nodes needed for a task, thereby
    reducing unnecessary context loading and reducing unsafe broad activation.

    Example route modes include direct expert route, debate route, quorum route,
    stop-signal route, contract-net auction route, emergency policy route, and
    federated aggregation route.

    ### Multi-Plane Memory And Engram Layer

    The memory system 122 separates source canon, compact canon, post-book
    addenda, assimilation ledger entries, project documents, trace ledgers,
    skill registries, expert genome registries, evaluation histories, provider
    scorecards, sandbox artifacts, and federated aggregate lessons. A retrieved
    memory carries provenance, status, freshness, and permission labels. The
    system can perform exact reference lookup, structured metadata lookup,
    semantic retrieval, graph-neighborhood lookup, and later hashed engram-style
    hot-memory lookup for frequent canon or workflow patterns.

    ### Recurrent Deliberation And Exit Gates

    The recurrent deliberation loop 126 updates internal hive state before
    external action. In each loop, the system may focus activations, route
    sparse experts, retrieve memory, run critique, update the HiveBlackboard,
    and evaluate risk and completeness. Exit gates can include confidence above
    threshold, risk below threshold, required policy gates passed, expert
    disagreement resolved, memory sufficient, maximum loops reached, operator
    checkpoint required, or policy forced stop.

    ### Sandboxed Self-Assimilation

    Referring to FIG. 4, a candidate capability can enter from operator prompt,
    research monitor, repository, paper, video, product, or trace discovery.
    The system pins source identity, reviews license and privacy class, extracts
    traits into a capability genome, runs a closed sandbox test, executes
    evaluations, applies immune/governance checks, writes a curator report, and
    promotes, rejects, blocks, or side-bars the candidate. The candidate cannot
    silently modify production memory, expert definitions, tools, prompts, or
    runtime state.

    ### Skill Systems And AFK Sandbox Agent Factory

    In some embodiments, focused skills are composed by an orchestrator into a
    skill system. A skill system can include transcript extraction, research,
    planning, implementation, review, packaging, scheduling, or bridge actions.
    A sandbox agent factory can launch planner, implementer, reviewer, and
    merger agents in isolated workspaces or containers. Merge-back is permitted
    only after policy scan, test evidence, checkpoint reference, and operator
    approval when required.

    ### Federated Learning Boundary

    Referring to FIG. 6, each local NexusNet deployment can produce local
    scorecards and deltas. A redaction and consent gate removes raw private
    content, secrets, proprietary source code, and unapproved transcripts. An
    aggregator receives approved summaries, computes aggregate lessons, and
    returns signed update candidates. The receiving deployment treats returned
    updates as candidates requiring sandbox and evaluation before promotion.

    ### VisualOps Control Panel

    The VisualOps control panel 136 renders brain path traces, active assistant
    orchestrators, expert capsules, route decisions, memory lookups, neural bus
    traffic, sandbox outcomes, evaluation scores, policy blocks, checkpoint
    references, federation state, and product readiness. Missing telemetry is
    shown as unavailable rather than fabricated.

    ### Technical Effects

    The described architecture can provide technical effects including reduced
    uncontrolled tool execution, reduced unnecessary expert activation,
    improved traceability of AI decisions, reversible autonomous update
    attempts, privacy-preserving cross-deployment learning, improved isolation
    of candidate capabilities, and improved operator visibility into AI system
    state.

    ## Example Embodiments

    1. A local-first desktop deployment uses local memory and project-root
       ledgers while routing selected tasks to hosted or local models.
    2. A coding deployment uses the sandbox agent factory to implement backlog
       tasks in isolated branches and requires reviewer and policy-gate evidence
       before merge.
    3. A research deployment monitors public sources, proposes candidate
       assimilations, and side-bars non-useful attempts as historical evidence.
    4. An enterprise deployment disables raw federated exports and shares only
       approved route-score summaries.
    5. A native-growth deployment uses traces, expert scorecards, dream outputs,
       and evaluation data as training material for later MoE-style model
       distillation.
""")


CLAIMS = wrap("""
    # Claims Draft

    The following claims are a technical drafting starting point only. Have a
    registered patent practitioner review and revise before filing if possible.

    ## Claim 1 - System

    1. A computer-implemented artificial intelligence harness system comprising:
       one or more processors; memory storing instructions executable by the one
       or more processors; a neural core service configured to mediate task
       execution; a hive node registry storing typed records for a plurality of
       harness nodes including assistant orchestrators, expert capsules, memory
       banks, model adapters, tool adapters, evaluators, sandbox runners, and
       policy gates; a neural bus configured to carry typed hive activations
       between the harness nodes; a cortex router configured to select a sparse
       subset of the harness nodes for a task based on at least capability,
       confidence, risk, privacy, cost, latency, permission, and historical
       reliability scores; a multi-plane memory system configured to return
       memory items with provenance and status labels; a recurrent deliberation
       loop configured to update internal harness state through one or more
       routing, memory, expert, critique, and policy iterations before an
       external action is permitted; an immune governance kernel configured to
       block a route or candidate update that lacks a required permission,
       provenance, sandbox, evaluation, or checkpoint condition; a
       checkpoint/rewind ledger configured to store pre-action rollback
       metadata; and an action/output subsystem configured to emit an external
       output only after the neural core service determines that required gates
       have passed.

    2. The system of claim 1, wherein the hive node registry stores, for each
       harness node, a node type, capabilities, allowed tools, write scope,
       privacy scope, concurrency safety, memory references, certification
       state, quarantine state, scorecard references, and a genome reference.

    3. The system of claim 1, wherein the typed hive activation includes an
       intent, task vector, risk vector, capability vector, memory references,
       policy labels, confidence, novelty, privacy class, and trace references.

    4. The system of claim 1, wherein the recurrent deliberation loop exits in
       response to at least one of a confidence threshold, a risk threshold, a
       policy-gate result, a disagreement-resolution result, a memory-sufficiency
       result, a maximum loop count, or an operator-checkpoint requirement.

    5. The system of claim 1, wherein the multi-plane memory system separates at
       least source canon records, post-book addendum records, assimilation
       ledger records, trace records, expert genome records, evaluation records,
       sandbox artifacts, and federated aggregate lessons.

    6. The system of claim 1, wherein the cortex router is configured to perform
       one or more of a direct expert route, debate route, quorum route,
       stop-signal route, contract-net auction route, emergency policy route, or
       federated aggregation route.

    7. The system of claim 1, further comprising a VisualOps control panel
       configured to display brain path traces, selected nodes, rejected nodes,
       memory lookups, policy blocks, sandbox outcomes, evaluation scores,
       checkpoint references, and federation status.

    8. The system of claim 1, further comprising a skill-system orchestrator
       configured to compose multiple focused skills into a workflow having
       typed handoffs, context limits, checkpoint locations, and output
       artifacts.

    9. The system of claim 1, further comprising a sandbox agent factory
       configured to run planner, implementer, reviewer, and merger stages in
       isolated execution environments and to permit merge-back only after
       policy, test, review, and checkpoint requirements are satisfied.

    10. The system of claim 1, wherein a candidate capability is promotable only
        after source identity pinning, license classification, privacy
        classification, sandbox testing, baseline comparison, evaluation, and
        immune-governance review.

    11. The system of claim 1, wherein the checkpoint/rewind ledger stores at
        least pre-write file state, active prompt state, memory reference state,
        selected route state, sandbox artifact state, candidate genome state,
        and promotion decision state.

    12. The system of claim 1, further comprising a federated learning plane
        configured to export approved, redacted deltas comprising one or more of
        route-score summaries, expert-performance scorecards, evaluation
        summaries, runtime-performance scorecards, or failure signatures while
        blocking export of raw private prompts, secrets, local file contents,
        proprietary code, and unredacted transcripts.

    13. The system of claim 1, wherein the immune governance kernel includes a
        plan-mode write jail, a tool execution registry, provider circuit
        breakers, prompt overlay controls, redaction filters, license filters,
        quarantine states, and human checkpoint gates.

    14. The system of claim 1, wherein the neural core service is configured to
        prevent user-interface actions, model providers, tool adapters, protocol
        adapters, or training jobs from bypassing brain-mediated routing and
        policy checks.

    ## Claim 15 - Method

    15. A computer-implemented method comprising: receiving a task or candidate
        capability; generating a typed hive activation from the task or
        candidate capability; attaching temporal lineage and provenance to the
        typed hive activation; retrieving memory items from a multi-plane memory
        system, wherein each retrieved memory item includes a provenance label
        and a status label; selecting, by a cortex router, a sparse subset of
        harness nodes from a hive node registry; executing one or more recurrent
        deliberation loops including expert computation, memory lookup, critique,
        and policy evaluation; determining, by an immune governance kernel,
        whether required gates have passed; creating a checkpoint reference
        before an external write or promotion action; and emitting an output or
        candidate disposition through a neural core service.

    16. The method of claim 15, further comprising rejecting or side-barring the
        candidate capability when sandbox or evaluation evidence fails a
        threshold.

    17. The method of claim 15, further comprising updating route weights or
        expert scorecards based on test results, regression results, user
        correction, policy violations, latency, cost, confidence, or usefulness.

    18. The method of claim 15, further comprising displaying, in a VisualOps
        interface, a trace of selected harness nodes, rejected harness nodes,
        memory operations, policy decisions, sandbox results, checkpoint
        references, and output metadata.

    19. The method of claim 15, further comprising composing a workflow from a
        plurality of focused skills by an orchestrator skill that passes typed
        outputs from one skill as inputs to another skill.

    20. The method of claim 15, further comprising receiving a federated update
        candidate, verifying redaction and consent metadata, and testing the
        federated update candidate in a sandbox before promotion.

    21. The method of claim 15, wherein the sparse subset includes at least one
        assistant orchestrator and at least one expert capsule.

    22. The method of claim 15, wherein the output comprises code, a document, a
        tool call, an application programming interface response, a visual state
        update, a training artifact, or a candidate assimilation report.

    ## Claim 23 - Computer-Readable Medium

    23. One or more non-transitory computer-readable storage media storing
        instructions that, when executed by one or more processors, cause the
        one or more processors to perform operations comprising: maintaining a
        hive node registry of typed artificial-intelligence harness nodes;
        routing typed hive activations over a neural bus; selecting sparse
        harness nodes through a cortex router; retrieving provenance-labeled
        memory from a multi-plane memory system; performing recurrent
        deliberation before external action; enforcing immune-governance gates;
        recording checkpoint and rewind metadata; and emitting an output through
        a brain-mediated action subsystem.

    24. The media of claim 23, wherein the operations further comprise
        classifying a candidate capability as promoted, rejected, blocked, or
        side-barred based on source, license, privacy, sandbox, evaluation,
        checkpoint, and policy evidence.

    25. The media of claim 23, wherein the operations further comprise
        generating training records for a later mixture-of-experts artificial
        intelligence model from trace records, expert scorecards, evaluation
        results, dream simulation outputs, and federated aggregate lessons.
""")


ABSTRACT = wrap("""
    # Abstract Draft

    A computer-implemented neural network harness for artificial intelligence
    includes a neural core service, a hive node registry, a neural bus, a cortex
    router, assistant orchestrators, expert capsules, memory banks, evaluators,
    sandbox runners, policy gates, a recurrent deliberation loop, an
    immune/governance kernel, a checkpoint/rewind ledger, and an action/output
    subsystem. The harness converts a task or candidate capability into a typed
    hive activation, retrieves provenance-labeled memory, selects a sparse set
    of harness nodes according to capability, confidence, risk, privacy, cost,
    latency, permission, and reliability scores, performs recurrent internal
    deliberation, and emits an output only after required gates pass. Candidate
    improvements are source-pinned, license-checked, privacy-classified,
    sandbox-tested, evaluated, and promoted, rejected, blocked, or side-barred
    with ledger evidence. A federated learning plane shares approved redacted
    score deltas without exporting raw private data.
""")


DRAWING_NOTES = wrap("""
    # Drawing Brief Description And Reference Numerals

    These drawings are a starting point for formal patent drawings. Review
    current USPTO drawing requirements before filing.

    ## Figure List

    - FIG. 1: Brain-first neural network harness system.
    - FIG. 2: Neural-harness plane stack.
    - FIG. 3: Task forward pass.
    - FIG. 4: Sandboxed self-assimilation loop.
    - FIG. 5: Skill-system and AFK sandbox agent factory.
    - FIG. 6: Privacy-preserving federated learning.
    - FIG. 7: Trace, checkpoint, and rewind ledgers.

    ## Reference Numerals

    | Numeral | Element |
    | --- | --- |
    | 100 | NexusNet neural network harness system |
    | 102 | Operator or workflow input |
    | 104 | Sensory/input plane |
    | 106 | Representation plane |
    | 108 | Temporal lineage ledger |
    | 110 | Neural bus |
    | 112 | Cortex router |
    | 114 | Neural core / NexusBrain |
    | 116 | Assistant orchestrator hive |
    | 118 | Expert capsule hive |
    | 120 | Mini-NexusNet specialist brain |
    | 122 | Multi-plane memory / engram layer |
    | 124 | HiveBlackboard residual state |
    | 126 | Recurrent deliberation loop |
    | 128 | Immune/governance kernel |
    | 130 | Sandbox and evaluation subsystem |
    | 132 | Checkpoint/rewind ledger |
    | 134 | Action/output subsystem |
    | 136 | VisualOps control panel |
    | 138 | Federated learning plane |
    | 140 | Ivy-League School optimizer |
    | 142 | Hive Curator AO |
    | 144 | Provider/tool/protocol adapters |
    | 146 | Source canon/addendum/assimilation ledger |
    | 150 | Hive activation |
    | 152 | Route decision |
    | 154 | Expert genome |
    | 156 | Hive trace |
    | 158 | Candidate capability |
    | 160 | Redacted federated delta |
""")


PRIOR_ART = wrap("""
    # Prior Art And Differentiator Worksheet

    This worksheet helps you organize known references for duty-of-disclosure
    review and claim differentiation. It is not a legal patentability opinion.

    ## Known Categories To Search

    - Multi-agent orchestration frameworks.
    - RAG and memory operating systems.
    - Mixture-of-experts routing.
    - Autonomous coding-agent sandboxes.
    - AI tool permission systems.
    - Federated learning systems.
    - AI eval and trace frameworks.
    - Agent UI/control panels.
    - Model self-improvement and curriculum systems.

    ## Differentiator Map

    | Conventional approach | NexusNet differentiator |
    | --- | --- |
    | Chatbot or wrapper forwards prompts to a model | Brain-mediated path with no-bypass policy, route trace, memory, critique, and output metadata |
    | RAG adds documents to prompt context | Multi-plane memory with provenance, canon/addendum status, temporal truth, and promotion gates |
    | Multi-agent framework delegates tasks | Neural-harness substrate maps agents/tools/models/memory to nodes, edges, weights, activations, recurrent loops, loss, and optimization |
    | Coding agent uses broad permissions | Sandboxed planner/implementer/reviewer/merger lanes with checkpoint, policy, tests, and rollback |
    | MoE model routes tokens among experts | Harness-level sparse routing activates AOs, expert capsules, tools, memory, models, sandboxes, and evaluators |
    | Self-improvement installs updates | Source-pinned candidate assimilation with license, privacy, sandbox, eval, immune, curator, and sidebar outcomes |
    | Federated learning shares model updates | Redacted route/expert/eval/runtime deltas with consent and no raw private data |
    | Dashboard shows status | VisualOps renders brain path, expert activations, memory operations, gates, disagreements, traces, and readiness |

    ## References To Consider For IDS Review

    Add exact publications, repositories, products, videos, papers, and dates.

    | Ref. no. | Reference | Date | Why material? | Submitted in IDS? |
    | --- | --- | --- | --- | --- |
    | R1 | [INSERT] | [INSERT] | [INSERT] | [YES/NO] |
""")


FILING = wrap("""
    # Filing Checklist And Patent Center Steps

    ## Decide Filing Type

    ### Option 1: Provisional First

    Use when you want a lower-cost priority placeholder while continuing to
    refine claims. A provisional does not issue as a patent and must be followed
    by a nonprovisional within the applicable deadline to benefit from the
    provisional filing date.

    Packet materials to use:
    - `03_INVENTION_DISCLOSURE_RECORD.md`
    - `04_NONPROVISIONAL_SPECIFICATION_DRAFT.md` converted into provisional
      disclosure form if desired
    - `drawings/NEXUSNET_PATENT_DRAWINGS.pdf`
    - Official provisional cover sheet
    - Micro-entity certification if applicable
    - Current provisional filing fee

    ### Option 2: Nonprovisional Utility Now

    Use when you are ready to file claims and begin examination.

    Packet materials to use:
    - `docx/NEXUSNET_UTILITY_SPECIFICATION_CLAIMS_ABSTRACT_DRAFT.docx`
    - `drawings/NEXUSNET_PATENT_DRAWINGS.pdf`
    - Official Application Data Sheet
    - Official inventor declaration/oath or later declaration strategy
    - Official micro-entity certification if applicable
    - Current filing, search, and examination fees

    ## Patent Center Steps

    1. Create or sign into a USPTO account.
    2. Open Patent Center.
    3. Start a new utility or provisional application.
    4. Enter bibliographic data from `02_USPTO_FORM_FIELD_WORKSHEETS.md`.
    5. Upload DOCX specification/claims/abstract material if filing
       nonprovisional.
    6. Upload drawings PDF.
    7. Attach official ADS, declaration, micro-entity certification, and any
       other required forms.
    8. Review validation warnings carefully.
    9. Pay fees using current micro-entity amounts only if eligibility is
       confirmed.
    10. Save the electronic acknowledgement receipt, application number,
        confirmation number, and filing date into this packet.

    ## Post-Filing Docket

    - Filing date:
    - Application number:
    - Confirmation number:
    - Customer number, if any:
    - Entity status claimed:
    - Nonprovisional deadline if provisional:
    - First Office action deadline:
    - IDS deadline review:
""")


IDS = wrap("""
    # Information Disclosure Statement Worksheet

    This is a worksheet only. If filing an IDS, use current official USPTO IDS
    forms and rules. A duty of disclosure can apply to information material to
    patentability.

    ## Known Technical References

    | No. | Reference | Type | Date | Material point | Include? |
    | --- | --- | --- | --- | --- | --- |
    | 1 | Attention Is All You Need | Paper | 2017 | Transformer attention and feed-forward architecture | [REVIEW] |
    | 2 | Sparsely-Gated Mixture-of-Experts | Paper | 2017 | Sparse expert activation | [REVIEW] |
    | 3 | Switch Transformer / GShard | Paper | 2020-2021 | Large sparse model routing | [REVIEW] |
    | 4 | Message Passing Neural Networks | Paper | 2017 | Graph message passing | [REVIEW] |
    | 5 | Graph Attention Networks | Paper | 2017 | Attention over graph neighbors | [REVIEW] |
    | 6 | Differentiable Neural Computers | Paper/blog | 2016 | External memory controller | [REVIEW] |
    | 7 | Retrieval-Augmented Generation | Paper | 2020 | Retrieval-supported generation | [REVIEW] |
    | 8 | Federated Averaging | Paper | 2017 | Privacy-preserving distributed learning | [REVIEW] |
    | 9 | Sandcastle-style sandbox coding agents | Repository/product | 2026 | Isolated agent execution pattern | [REVIEW] |
    | 10 | Agent skill systems / Claude Code skills | Product docs/videos | 2026 | Composable skill workflows | [REVIEW] |

    ## Product/Repository References

    | No. | Product/repo | URL | Date reviewed | Potential relevance |
    | --- | --- | --- | --- | --- |
    | P1 | [INSERT] | [INSERT] | [INSERT] | [INSERT] |

    ## Disclosure Decision Notes

    Record why each reference is or is not included in a formal IDS.
""")


SUPPORT_CHART = wrap("""
    # Source Canon To Claim Support Chart

    Use this chart to check whether each proposed claim element is supported by
    the NexusNet source record before filing. This is not a patentability
    opinion. It is a written-description and drafting-support aid.

    ## Source Records

    - S1: `docs/NEXUSNET_COMPLETE_CHAT_CANON_BOOK_2026-04-28.md`
    - S2: `docs/NEXUSNET_CHAT_RESEARCH_CANON_2026-04-28.md`
    - S3: `docs/NEXUSNET_POST_BOOK_CANON_ADDENDUM_2026-05-01.md`
    - S4: `docs/superpowers/specs/2026-05-01-nexusnet-hive-neural-substrate-v0-design.md`
    - S5: `nexusnet/core/brain.py`
    - S6: `nexusnet/canon/realization.py`

    ## Claim Element Support

    | Claim element | Draft claim refs | Primary support | Notes to verify |
    | --- | --- | --- | --- |
    | Brain-mediated neural core service | 1, 14, 15, 23 | S1, S2, S5, S6 | Confirm every user-facing path is intended to route through NexusBrain or equivalent. |
    | Hive node registry | 1, 2, 23 | S2, S3, S4, S6 | Confirm final registry fields and expert/AO node taxonomy. |
    | Neural bus carrying typed activations | 1, 3, 15, 23 | S2, S4 | Confirm v0 data contract names. |
    | Sparse cortex router / MoE-style selection | 1, 6, 15, 23 | S1, S2, S4, S6 | Confirm selected scoring dimensions. |
    | Multi-plane memory / engram layer | 1, 5, 15, 23 | S1, S2, S4 | Confirm memory plane taxonomy and provenance/status labels. |
    | Recurrent deliberation loop and exit gates | 1, 4, 15, 23 | S3, S4 | Confirm loop metadata and exit gate fields. |
    | Immune/governance kernel | 1, 10, 13, 15, 23 | S2, S3, S4 | Confirm no-bypass policy and write gates. |
    | Checkpoint/rewind ledger | 1, 11, 15, 23 | S3, S4 | Confirm pre-write snapshot and rollback metadata. |
    | Skill-system orchestrator | 8, 19 | S3, S4 | Confirm focused skills plus orchestrator, not mega-skill pattern. |
    | Sandbox agent factory | 9 | S3, S4 | Confirm planner/implementer/reviewer/merger lane evidence. |
    | Candidate assimilation and sidebar | 10, 16, 24 | S2, S3, S4 | Confirm sandbox/eval/promotion/rejection ledger states. |
    | Federated learning plane | 12, 20, 25 | S1, S2, S3, S4 | Confirm privacy boundary and redacted delta types. |
    | VisualOps cockpit | 7, 18 | S1, S2, S4 | Confirm live surfaces and missing telemetry behavior. |
    | Native MoE evolution path | 25 | S1, S2, S3, S4 | Confirm this is future growth path, not current implementation claim. |

    ## Filing Review Notes

    - Remove or narrow any claim element that cannot be enabled in the
      specification.
    - Avoid claiming only a result, such as "self-improves," without concrete
      data structures, routing steps, gates, and ledgers.
    - Separate implemented features from planned embodiments when discussing
      commercial readiness.
""")


PROVISIONAL = wrap("""
    # Provisional Application Disclosure Draft

    This document is a shorter provisional-style disclosure assembled from the
    nonprovisional draft. A provisional application can be useful when you want
    an earlier filing date while continuing to refine claims, drawings, and
    implementation details. A provisional does not become a patent by itself.

    ## Title

    Hive Mind Harness for Artificial Intelligence with Governed Sparse Routing

    ## Inventor

    [INSERT FULL LEGAL NAME]

    ## Technical Field

    Artificial intelligence harnesses, multi-agent orchestration, model routing,
    memory-augmented computing, sandboxed autonomous software updates,
    federated learning, and AI control-panel systems.

    ## Problem

    AI applications commonly combine model providers, retrieval, tools, agents,
    and dashboards without a unified technical substrate that governs routing,
    memory, permissions, self-improvement, rollback, evaluation, and operator
    visibility. Such systems can be difficult to audit and can permit unsafe or
    poorly evidenced changes.

    ## Proposed Solution

    NexusNet provides a brain-first neural network harness. A canonical neural
    core receives or mediates tasks, creates typed hive activations, retrieves
    provenance-labeled memory, routes activations through a neural bus, selects
    sparse assistant orchestrators and expert capsules through a cortex router,
    performs recurrent deliberation, applies immune/governance gates, records
    checkpoint and rewind metadata, and emits outputs only after required gates
    pass.

    ## Main Components

    - Neural core / NexusBrain.
    - Assistant orchestrator hive.
    - Expert capsule hive and Mini-NexusNets.
    - Neural bus and HiveBlackboard.
    - Cortex router with sparse MoE-style node selection.
    - Multi-plane memory and engram layer.
    - Recurrent deliberation loop.
    - Sandbox and evaluation subsystem.
    - Immune/governance kernel.
    - Checkpoint/rewind ledger.
    - Skill-system orchestrator.
    - Sandbox agent factory.
    - Federated learning plane.
    - Ivy-League School optimizer.
    - Hive Curator AO.
    - VisualOps control panel.

    ## Operation

    1. Receive operator task, workflow input, candidate update, tool result, or
       federated signal.
    2. Assign provenance, trust, privacy, and policy labels.
    3. Convert the signal into a typed activation.
    4. Retrieve relevant canon, memory, evidence, traces, and candidate records.
    5. Focus active signals and select sparse experts, AOs, tools, memory banks,
       models, sandboxes, and evaluators.
    6. Run recurrent deliberation until an exit gate is satisfied.
    7. Apply immune/governance checks.
    8. Create checkpoint and rollback references before writes or promotions.
    9. Emit output, promote candidate, reject candidate, block candidate, or
       side-bar candidate with evidence.
    10. Feed scorecards to curator, training, routing, and federated learning
        planes.

    ## Advantages

    The system improves technical control over AI operations by making routing,
    memory use, tool execution, self-improvement, evaluation, and rollback
    explicit, logged, inspectable, and governed. The architecture can reduce
    unnecessary broad activation, reduce unsafe autonomous writes, preserve
    provenance, support privacy-preserving learning across deployments, and
    generate training evidence for later native MoE-style model development.

    ## Drawings

    Attach `drawings/NEXUSNET_PATENT_DRAWINGS.pdf` to the provisional filing if
    filing this disclosure as a provisional application.
""")


TRANSMITTAL = wrap("""
    # Transmittal Letter Template

    [DATE]

    Commissioner for Patents
    P.O. Box 1450
    Alexandria, VA 22313-1450

    Re: Patent Application for "Hive Mind Harness for Artificial Intelligence
    with Governed Sparse Routing"

    Applicant/Inventor: [INSERT]
    Attorney Docket No.: NEXUSNET-001-US

    Dear Commissioner:

    Submitted herewith for filing is a patent application for the above-titled
    invention. The submission includes the specification, claims, abstract,
    drawings, and applicable filing documents. Applicant requests recognition
    of micro-entity status if the attached certification is complete and
    accurate.

    Please charge any required fees as authorized in Patent Center or as
    otherwise provided by the applicant.

    Respectfully submitted,

    [SIGNATURE]
    [NAME]
    [ADDRESS]
    [PHONE]
    [EMAIL]
""")


MICRO_NOTE = wrap("""
    # Micro-Entity Cover Note Template

    Re: Micro-Entity Certification

    Applicant believes micro-entity status applies based on the completed
    certification submitted with this application. Applicant has reviewed the
    small entity requirement, prior application limitation, income threshold,
    and assignment/obligation limitation, and understands that entity status
    must be corrected if the facts change.

    This note is not a substitute for the official USPTO certification form.
""")


DECL_NOTE = wrap("""
    # Inventor Declaration Preparation Note

    Each inventor should review the final application before signing any oath
    or declaration. Confirm that:

    - The named inventors are correct.
    - The claims are supported by the disclosure.
    - Each inventor believes they are an original inventor of the claimed
      subject matter.
    - The duty to disclose material information is understood.
    - No signature is added to an unofficial worksheet by mistake.
""")


def build_markdown_files() -> None:
    files = {
        "READ_ME_FIRST_READ_ORDER.md": READ_ORDER,
        "README_PRO_SE_PACKET.md": README,
        "00_USE_AND_LIMITS.md": USE_LIMITS,
        "01_MICRO_ENTITY_WORKSHEET.md": MICRO_ENTITY,
        "02_USPTO_FORM_FIELD_WORKSHEETS.md": FORMS,
        "03_INVENTION_DISCLOSURE_RECORD.md": DISCLOSURE,
        "04_NONPROVISIONAL_SPECIFICATION_DRAFT.md": SPEC,
        "05_CLAIMS_DRAFT.md": CLAIMS,
        "06_ABSTRACT_DRAFT.md": ABSTRACT,
        "07_DRAWING_BRIEF_DESCRIPTION_AND_REFERENCE_NUMERALS.md": DRAWING_NOTES,
        "08_PRIOR_ART_AND_DIFFERENTIATOR_WORKSHEET.md": PRIOR_ART,
        "09_FILING_CHECKLIST_AND_PATENT_CENTER_STEPS.md": FILING,
        "10_IDS_WORKSHEET.md": IDS,
        "11_SOURCE_CANON_TO_CLAIM_SUPPORT_CHART.md": SUPPORT_CHART,
        "12_PROVISIONAL_APPLICATION_DISCLOSURE_DRAFT.md": PROVISIONAL,
        "letters/TRANSMITTAL_LETTER_TEMPLATE.md": TRANSMITTAL,
        "letters/MICRO_ENTITY_COVER_NOTE.md": MICRO_NOTE,
        "letters/INVENTOR_DECLARATION_PREP_NOTE.md": DECL_NOTE,
    }
    for rel, content in files.items():
        write(OUT / rel, content)


def add_md_to_doc(doc: Document, content: str) -> None:
    for raw in content.splitlines():
        line = raw.rstrip()
        if not line:
            doc.add_paragraph()
            continue
        if line.startswith("# "):
            add_heading(doc, line[2:], 1)
        elif line.startswith("## "):
            add_heading(doc, line[3:], 2)
        elif line.startswith("### "):
            add_heading(doc, line[4:], 3)
        elif line.startswith("- "):
            add_bullet(doc, line[2:])
        elif line.startswith("    "):
            add_para(doc, line.strip())
        else:
            add_para(doc, line)


def build_utility_docx() -> None:
    DOCX_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    style_doc(doc)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(TITLE)
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(31, 78, 121)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run("Draft utility specification, claims, and abstract for pro se review").italic = True
    doc.add_paragraph(f"Generated: {DATE}")
    doc.add_paragraph("Attorney Docket No.: NEXUSNET-001-US")
    doc.add_paragraph("Inventor/Applicant: [INSERT FULL LEGAL NAME]")
    doc.add_page_break()
    add_md_to_doc(doc, SPEC)
    doc.add_page_break()
    add_md_to_doc(doc, CLAIMS)
    doc.add_page_break()
    add_md_to_doc(doc, ABSTRACT)
    path = DOCX_DIR / "NEXUSNET_UTILITY_SPECIFICATION_CLAIMS_ABSTRACT_DRAFT.docx"
    doc.save(path)


def build_forms_docx() -> None:
    doc = Document()
    style_doc(doc)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("NexusNet Pro Se Filing Worksheets")
    r.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = RGBColor(31, 78, 121)
    doc.add_paragraph(f"Generated: {DATE}")
    doc.add_paragraph("Use these worksheets to complete current official USPTO forms.")
    doc.add_page_break()
    for content in [MICRO_ENTITY, FORMS, DISCLOSURE, PRIOR_ART, FILING, IDS, SUPPORT_CHART, PROVISIONAL, TRANSMITTAL, MICRO_NOTE, DECL_NOTE]:
        add_md_to_doc(doc, content)
        doc.add_page_break()
    doc.save(DOCX_DIR / "NEXUSNET_PRO_SE_FORMS_AND_WORKSHEETS.docx")


def box(c: canvas.Canvas, x: float, y: float, w: float, h: float, text: str, ref: str | None = None) -> None:
    c.setStrokeColor(colors.black)
    c.setLineWidth(1)
    c.roundRect(x, y, w, h, 6, stroke=1, fill=0)
    c.setFont("Helvetica", 8)
    lines = textwrap.wrap(text, width=max(10, int(w / 5.5)))
    ty = y + h - 14
    for line in lines[:5]:
        c.drawCentredString(x + w / 2, ty, line)
        ty -= 10
    if ref:
        c.setFont("Helvetica-Bold", 8)
        c.drawRightString(x + w - 4, y + 5, ref)


def arrow(c: canvas.Canvas, x1: float, y1: float, x2: float, y2: float) -> None:
    c.setLineWidth(0.8)
    c.line(x1, y1, x2, y2)
    dx = x2 - x1
    dy = y2 - y1
    if abs(dx) > abs(dy):
        sign = 1 if dx >= 0 else -1
        c.line(x2, y2, x2 - 6 * sign, y2 + 3)
        c.line(x2, y2, x2 - 6 * sign, y2 - 3)
    else:
        sign = 1 if dy >= 0 else -1
        c.line(x2, y2, x2 + 3, y2 - 6 * sign)
        c.line(x2, y2, x2 - 3, y2 - 6 * sign)


def start_sheet(c: canvas.Canvas, fig: str, title: str) -> None:
    c.setPageSize(letter)
    c.setFont("Helvetica-Bold", 11)
    c.drawCentredString(4.25 * inch, 10.45 * inch, f"{fig} - {title}")
    c.setLineWidth(0.5)
    c.rect(0.75 * inch, 0.5 * inch, 7.0 * inch, 9.75 * inch, stroke=1, fill=0)


def fig1(c: canvas.Canvas) -> None:
    start_sheet(c, "FIG. 1", "Brain-First Neural Network Harness System 100")
    y = 9.35 * inch
    box(c, 0.95 * inch, y, 1.25 * inch, 0.42 * inch, "Operator input", "102")
    box(c, 2.45 * inch, y, 1.25 * inch, 0.42 * inch, "Sensory / input", "104")
    box(c, 3.95 * inch, y, 1.25 * inch, 0.42 * inch, "Representation", "106")
    box(c, 5.45 * inch, y, 1.25 * inch, 0.42 * inch, "Temporal lineage", "108")
    arrow(c, 2.2 * inch, y + 0.21 * inch, 2.45 * inch, y + 0.21 * inch)
    arrow(c, 3.7 * inch, y + 0.21 * inch, 3.95 * inch, y + 0.21 * inch)
    arrow(c, 5.2 * inch, y + 0.21 * inch, 5.45 * inch, y + 0.21 * inch)
    box(c, 3.0 * inch, 8.15 * inch, 2.1 * inch, 0.55 * inch, "Neural bus", "110")
    arrow(c, 6.1 * inch, y, 4.05 * inch, 8.7 * inch)
    box(c, 1.0 * inch, 6.85 * inch, 1.55 * inch, 0.62 * inch, "AO hive", "116")
    box(c, 3.05 * inch, 6.85 * inch, 1.55 * inch, 0.62 * inch, "NexusBrain / neural core", "114")
    box(c, 5.1 * inch, 6.85 * inch, 1.55 * inch, 0.62 * inch, "Cortex router", "112")
    arrow(c, 3.65 * inch, 8.15 * inch, 1.75 * inch, 7.47 * inch)
    arrow(c, 4.05 * inch, 8.15 * inch, 3.85 * inch, 7.47 * inch)
    arrow(c, 4.45 * inch, 8.15 * inch, 5.85 * inch, 7.47 * inch)
    box(c, 0.95 * inch, 5.55 * inch, 1.4 * inch, 0.62 * inch, "Expert capsule hive", "118")
    box(c, 2.55 * inch, 5.55 * inch, 1.4 * inch, 0.62 * inch, "Mini-NexusNet", "120")
    box(c, 4.15 * inch, 5.55 * inch, 1.4 * inch, 0.62 * inch, "Memory / engram", "122")
    box(c, 5.75 * inch, 5.55 * inch, 1.4 * inch, 0.62 * inch, "Adapters", "144")
    for x in [1.65, 3.25, 4.85, 6.45]:
        arrow(c, 5.85 * inch, 6.85 * inch, x * inch, 6.17 * inch)
    box(c, 1.0 * inch, 4.3 * inch, 1.55 * inch, 0.62 * inch, "HiveBlackboard", "124")
    box(c, 3.05 * inch, 4.3 * inch, 1.55 * inch, 0.62 * inch, "Recurrent loop", "126")
    box(c, 5.1 * inch, 4.3 * inch, 1.55 * inch, 0.62 * inch, "Immune / governance", "128")
    arrow(c, 3.85 * inch, 6.85 * inch, 3.85 * inch, 4.92 * inch)
    box(c, 1.0 * inch, 3.0 * inch, 1.55 * inch, 0.62 * inch, "Sandbox / eval", "130")
    box(c, 3.05 * inch, 3.0 * inch, 1.55 * inch, 0.62 * inch, "Checkpoint / rewind", "132")
    box(c, 5.1 * inch, 3.0 * inch, 1.55 * inch, 0.62 * inch, "Action / output", "134")
    arrow(c, 5.85 * inch, 4.3 * inch, 5.85 * inch, 3.62 * inch)
    box(c, 1.0 * inch, 1.75 * inch, 1.55 * inch, 0.62 * inch, "VisualOps", "136")
    box(c, 3.05 * inch, 1.75 * inch, 1.55 * inch, 0.62 * inch, "Federated plane", "138")
    box(c, 5.1 * inch, 1.75 * inch, 1.55 * inch, 0.62 * inch, "Ivy-League School", "140")


def fig2(c: canvas.Canvas) -> None:
    start_sheet(c, "FIG. 2", "Neural-Harness Plane Stack")
    planes = [
        ("Sensory/Input", "104"),
        ("Representation", "106"),
        ("Temporal/Positional", "108"),
        ("Neural Bus / Message Passing", "110"),
        ("Attention / Focus", "112A"),
        ("Sparse MoE Router", "112"),
        ("Expert Computation", "118"),
        ("Memory / Engram", "122"),
        ("Recurrent Deliberation", "126"),
        ("Learning / Eval / Loss", "130"),
        ("Optimizer / School", "140"),
        ("Federated Learning", "138"),
        ("Immune / Governance", "128"),
        ("Curator / Pruning", "142"),
        ("Action / Output", "134"),
        ("Checkpoint / Rewind", "132"),
    ]
    x = 1.2 * inch
    y = 9.3 * inch
    for i, (label, ref) in enumerate(planes):
        box(c, x, y - i * 0.5 * inch, 5.7 * inch, 0.32 * inch, label, ref)
        if i < len(planes) - 1:
            arrow(c, 4.05 * inch, y - i * 0.5 * inch, 4.05 * inch, y - (i + 1) * 0.5 * inch + 0.32 * inch)


def fig3(c: canvas.Canvas) -> None:
    start_sheet(c, "FIG. 3", "Task Forward Pass")
    steps = [
        ("Receive task", "102"),
        ("Create activation", "150"),
        ("Retrieve memory", "122"),
        ("Select sparse nodes", "152"),
        ("Loop deliberation", "126"),
        ("Run experts", "118"),
        ("Evaluate loss", "130"),
        ("Policy gate", "128"),
        ("Checkpoint", "132"),
        ("Emit output", "134"),
    ]
    y = 9.2 * inch
    for i, (label, ref) in enumerate(steps):
        x = 1.2 * inch if i % 2 == 0 else 4.25 * inch
        yy = y - (i // 2) * 1.35 * inch
        box(c, x, yy, 2.0 * inch, 0.55 * inch, label, ref)
        if i < len(steps) - 1:
            nx = 1.2 * inch if (i + 1) % 2 == 0 else 4.25 * inch
            ny = y - ((i + 1) // 2) * 1.35 * inch
            arrow(c, x + 2.0 * inch, yy + 0.27 * inch, nx, ny + 0.27 * inch)


def fig4(c: canvas.Canvas) -> None:
    start_sheet(c, "FIG. 4", "Sandboxed Self-Assimilation Loop")
    labels = [
        ("Candidate source", "158"),
        ("Source/license/privacy pin", "146"),
        ("Capability genome", "154"),
        ("Closed sandbox", "130"),
        ("Baseline eval", "130A"),
        ("Immune review", "128"),
        ("Curator report", "142"),
        ("Promote / reject / sidebar", "152"),
        ("Checkpoint and ledger", "132"),
    ]
    coords = [(1.0, 8.8), (3.3, 8.8), (5.6, 8.8), (5.6, 7.0), (3.3, 7.0), (1.0, 7.0), (1.0, 5.2), (3.3, 5.2), (5.6, 5.2)]
    for (label, ref), (x, y) in zip(labels, coords):
        box(c, x * inch, y * inch, 1.65 * inch, 0.58 * inch, label, ref)
    for i in range(len(coords) - 1):
        x1, y1 = coords[i]
        x2, y2 = coords[i + 1]
        arrow(c, (x1 + 1.65) * inch, (y1 + 0.29) * inch, x2 * inch, (y2 + 0.29) * inch)
    arrow(c, 6.4 * inch, 5.2 * inch, 6.4 * inch, 8.8 * inch)


def fig5(c: canvas.Canvas) -> None:
    start_sheet(c, "FIG. 5", "Skill-System And AFK Sandbox Agent Factory")
    box(c, 1.0 * inch, 8.8 * inch, 1.7 * inch, 0.55 * inch, "Orchestrator skill system", "170")
    labels = [("Research skill", "172"), ("Planning skill", "174"), ("Implementer agent", "176"), ("Reviewer agent", "178"), ("Merger gate", "180")]
    for i, (label, ref) in enumerate(labels):
        box(c, (1.0 + i * 1.25) * inch, 7.45 * inch, 1.05 * inch, 0.75 * inch, label, ref)
        if i > 0:
            arrow(c, (1.0 + (i - 1) * 1.25 + 1.05) * inch, 7.82 * inch, (1.0 + i * 1.25) * inch, 7.82 * inch)
    box(c, 1.0 * inch, 5.9 * inch, 1.7 * inch, 0.6 * inch, "Isolated worktree / container", "182")
    box(c, 3.25 * inch, 5.9 * inch, 1.7 * inch, 0.6 * inch, "Policy and tests", "184")
    box(c, 5.5 * inch, 5.9 * inch, 1.7 * inch, 0.6 * inch, "Checkpoint merge-back", "186")
    arrow(c, 2.7 * inch, 6.2 * inch, 3.25 * inch, 6.2 * inch)
    arrow(c, 4.95 * inch, 6.2 * inch, 5.5 * inch, 6.2 * inch)


def fig6(c: canvas.Canvas) -> None:
    start_sheet(c, "FIG. 6", "Privacy-Preserving Federated Learning")
    for i, x in enumerate([1.0, 3.2, 5.4], start=1):
        box(c, x * inch, 8.4 * inch, 1.35 * inch, 0.7 * inch, f"Local NexusNet node {i}", f"190{i}")
        box(c, x * inch, 7.25 * inch, 1.35 * inch, 0.6 * inch, "Redaction / consent gate", "192")
        arrow(c, (x + 0.67) * inch, 8.4 * inch, (x + 0.67) * inch, 7.85 * inch)
    box(c, 2.95 * inch, 5.75 * inch, 2.0 * inch, 0.7 * inch, "Federated aggregator", "194")
    for x in [1.67, 3.87, 6.07]:
        arrow(c, x * inch, 7.25 * inch, 3.95 * inch, 6.45 * inch)
    box(c, 2.95 * inch, 4.35 * inch, 2.0 * inch, 0.7 * inch, "Signed update candidate", "196")
    arrow(c, 3.95 * inch, 5.75 * inch, 3.95 * inch, 5.05 * inch)
    box(c, 2.95 * inch, 3.0 * inch, 2.0 * inch, 0.7 * inch, "Sandbox before local promotion", "130")
    arrow(c, 3.95 * inch, 4.35 * inch, 3.95 * inch, 3.7 * inch)


def fig7(c: canvas.Canvas) -> None:
    start_sheet(c, "FIG. 7", "Trace, Checkpoint, And Rewind Data Structures")
    tables = [
        ("HiveActivation 150", ["activation_id", "intent", "risk_vector", "memory_refs", "policy_labels", "privacy_class"]),
        ("RouteDecision 152", ["decision_id", "selected_nodes", "rejected_nodes", "scores", "required_gates", "exit_state"]),
        ("ExpertGenome 154", ["genome_id", "parents", "capability_traits", "tool_traits", "eval_traits", "mutation_history"]),
        ("HiveTrace 156", ["trace_id", "event_type", "node_refs", "policy_refs", "checkpoint_refs", "result"]),
    ]
    y = 8.9
    for title, fields in tables:
        box(c, 1.2 * inch, y * inch, 5.6 * inch, 0.38 * inch, title, None)
        c.setFont("Helvetica", 8)
        fy = (y - 0.27) * inch
        for f in fields:
            c.drawString(1.45 * inch, fy, f"- {f}")
            fy -= 0.16 * inch
        y -= 1.85


def build_drawings_pdf() -> None:
    DRAWINGS_DIR.mkdir(parents=True, exist_ok=True)
    pdf = DRAWINGS_DIR / "NEXUSNET_PATENT_DRAWINGS.pdf"
    c = canvas.Canvas(str(pdf), pagesize=letter)
    for fn in [fig1, fig2, fig3, fig4, fig5, fig6, fig7]:
        fn(c)
        c.showPage()
    c.save()


def build_simple_svg_files() -> None:
    DRAWINGS_DIR.mkdir(parents=True, exist_ok=True)
    svg_template = """<svg xmlns="http://www.w3.org/2000/svg" width="850" height="1100" viewBox="0 0 850 1100">
  <rect x="75" y="50" width="700" height="975" fill="white" stroke="black"/>
  <text x="425" y="90" text-anchor="middle" font-family="Arial" font-size="22" font-weight="bold">{title}</text>
  {body}
</svg>
"""
    bodies = {
        "FIG_1_SYSTEM_OVERVIEW.svg": '<rect x="300" y="390" width="250" height="80" fill="white" stroke="black"/><text x="425" y="435" text-anchor="middle" font-family="Arial" font-size="18">NexusBrain 114</text><rect x="110" y="250" width="180" height="70" fill="white" stroke="black"/><text x="200" y="292" text-anchor="middle" font-family="Arial" font-size="16">AO Hive 116</text><rect x="560" y="250" width="180" height="70" fill="white" stroke="black"/><text x="650" y="292" text-anchor="middle" font-family="Arial" font-size="16">Router 112</text><rect x="110" y="560" width="180" height="70" fill="white" stroke="black"/><text x="200" y="602" text-anchor="middle" font-family="Arial" font-size="16">Memory 122</text><rect x="560" y="560" width="180" height="70" fill="white" stroke="black"/><text x="650" y="602" text-anchor="middle" font-family="Arial" font-size="16">Governance 128</text>',
        "FIG_2_PLANE_STACK.svg": "".join(f'<rect x="150" y="{130+i*52}" width="550" height="36" fill="white" stroke="black"/><text x="425" y="{154+i*52}" text-anchor="middle" font-family="Arial" font-size="15">{label}</text>' for i, label in enumerate(["Input 104","Representation 106","Temporal 108","Neural Bus 110","Attention 112A","Sparse Router 112","Expert Compute 118","Memory 122","Recurrent Loop 126","Eval/Loss 130","School 140","Federation 138","Governance 128","Action 134","Checkpoint 132"])),
        "FIG_3_FORWARD_PASS.svg": '<text x="425" y="220" text-anchor="middle" font-family="Arial" font-size="18">Task -> Activation -> Memory -> Router -> Experts -> Eval -> Gate -> Output</text>',
        "FIG_4_ASSIMILATION_LOOP.svg": '<circle cx="425" cy="500" r="250" fill="white" stroke="black"/><text x="425" y="500" text-anchor="middle" font-family="Arial" font-size="18">Candidate Assimilation Loop 158</text>',
        "FIG_5_SKILL_AGENT_FACTORY.svg": '<text x="425" y="260" text-anchor="middle" font-family="Arial" font-size="18">Skill System Orchestrator 170</text><text x="425" y="360" text-anchor="middle" font-family="Arial" font-size="18">Planner -> Implementer -> Reviewer -> Merger</text>',
        "FIG_6_FEDERATED_LEARNING.svg": '<text x="425" y="260" text-anchor="middle" font-family="Arial" font-size="18">Local Nodes -> Redacted Deltas 160 -> Aggregator -> Sandbox Candidate</text>',
        "FIG_7_TRACE_LEDGER.svg": '<text x="425" y="260" text-anchor="middle" font-family="Arial" font-size="18">Activation 150 | Route 152 | Genome 154 | Trace 156 | Checkpoint 132</text>',
    }
    for name, body in bodies.items():
        title = name.replace("_", " ").replace(".svg", "")
        write(DRAWINGS_DIR / name, svg_template.format(title=title, body=body))


def build_manifest() -> None:
    paths = sorted(
        p.relative_to(OUT).as_posix()
        for p in OUT.rglob("*")
        if p.is_file() and "_qa_renders" not in p.relative_to(OUT).parts
    )
    manifest = "# Packet Manifest\n\n" + "\n".join(f"- `{p}`" for p in paths) + "\n"
    write(OUT / "PACKAGE_MANIFEST.md", manifest)


def main() -> None:
    for d in [OUT, DOCX_DIR, DRAWINGS_DIR, LETTERS_DIR, QA_DIR]:
        d.mkdir(parents=True, exist_ok=True)
    build_markdown_files()
    build_utility_docx()
    build_forms_docx()
    build_drawings_pdf()
    build_simple_svg_files()
    build_manifest()
    print(f"Generated patent packet at {OUT}")


if __name__ == "__main__":
    main()
